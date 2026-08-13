"""Builds an Excel (.xlsx) report from a job's results, reusing the
exact same data already served via the JSON API (JobResultsResponse /
LogEntry), so the export always matches what's shown in the UI.

Lives under api/ rather than core/ deliberately: it consumes the
pydantic response models built for the HTTP layer, which keeps core/
free of API-shaped types (see CONTRIBUTING.md).
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from imagecompare.models import ComparisonSummary, JobResultsResponse, LogEntry

_FONT_NAME = "Arial"
_HEADER_FONT = Font(name=_FONT_NAME, bold=True, color="FFFFFF", size=10)
_HEADER_FILL = PatternFill(start_color="22262C", end_color="22262C", fill_type="solid")
_TITLE_FONT = Font(name=_FONT_NAME, bold=True, size=14)
_LABEL_FONT = Font(name=_FONT_NAME, bold=True, size=10)
_BODY_FONT = Font(name=_FONT_NAME, size=10)
_PERCENT_FORMAT = '0.0"%"'

_SUMMARY_FIELD_LABELS: list[tuple[str, str]] = [
    ("mode", "Comparison mode"),
    ("folder1", "Folder A"),
    ("folder2", "Folder B"),
    ("total_images_a", "Images scanned in folder A"),
    ("total_images_b", "Images scanned in folder B"),
    ("total_pairs_compared", "Total pairs compared"),
    ("total_matches", "Matches found"),
    ("exact_duplicates", "Exact duplicates (MD5)"),
    ("highest_similarity", "Highest similarity"),
    ("average_similarity", "Average similarity"),
    ("lowest_similarity", "Lowest similarity"),
    ("visual_threshold", "Visual threshold used"),
    ("deep_rotation_used", "Deep Rotation used"),
    ("elapsed_seconds", "Elapsed time (seconds)"),
]

_MATCH_HEADERS = [
    "#",
    "Folder A - File Name",
    "Folder A - Relative Path",
    "Folder A - Full Path",
    "Folder A - MD5",
    "Folder B - File Name",
    "Folder B - Relative Path",
    "Folder B - Full Path",
    "Folder B - MD5",
    "Visual Similarity",
    "Filename Similarity",
    "Exact Duplicate (MD5)",
    "Best Rotation Angle (deg)",
]

_LOG_HEADERS = ["Timestamp", "Level", "Message"]


def build_report_workbook(
    results: JobResultsResponse,
    log_entries: list[LogEntry] | None = None,
) -> BytesIO:
    """Build the full report workbook and return it as an in-memory
    buffer (position reset to 0, ready to stream as a response body)."""
    wb = Workbook()

    summary_sheet = wb.active
    summary_sheet.title = "Summary"
    _write_summary_sheet(summary_sheet, results.summary)

    matches_sheet = wb.create_sheet("Matches")
    _write_matches_sheet(matches_sheet, results)

    if log_entries:
        log_sheet = wb.create_sheet("Activity Log")
        _write_log_sheet(log_sheet, log_entries)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _write_summary_sheet(ws: Worksheet, summary: ComparisonSummary | None) -> None:
    ws["A1"] = "Image Compare - Comparison Report"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:B1")

    if summary is None:
        ws["A3"] = "No summary available for this job."
        ws["A3"].font = _BODY_FONT
        ws.column_dimensions["A"].width = 40
        return

    data = summary.model_dump()
    row = 3
    for field_name, label in _SUMMARY_FIELD_LABELS:
        value = data.get(field_name)
        ws.cell(row=row, column=1, value=label).font = _LABEL_FONT
        cell = ws.cell(row=row, column=2, value=_format_summary_value(field_name, value))
        cell.font = _BODY_FONT
        if field_name in (
            "highest_similarity",
            "average_similarity",
            "lowest_similarity",
            "visual_threshold",
        ):
            cell.number_format = _PERCENT_FORMAT
        row += 1

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 60


def _format_summary_value(field_name: str, value: object) -> object:
    if value is None:
        return "-"
    if field_name == "deep_rotation_used":
        return "Yes" if value else "No"
    return value


def _write_matches_sheet(ws: Worksheet, results: JobResultsResponse) -> None:
    for col_idx, header in enumerate(_MATCH_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")

    for row_idx, m in enumerate(results.matches, start=2):
        values = [
            row_idx - 1,
            m.image1_name,
            m.image1_relative_path,
            m.image1_path,
            m.image1_md5,
            m.image2_name,
            m.image2_relative_path,
            m.image2_path,
            m.image2_md5,
            m.visual_similarity,
            m.filename_similarity,
            "Yes" if m.is_exact_duplicate else "No",
            m.best_rotation_angle if m.best_rotation_angle is not None else "",
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = _BODY_FONT
            if col_idx in (10, 11):  # similarity percentage columns
                cell.number_format = _PERCENT_FORMAT

    if results.matches:
        last_row = len(results.matches) + 1
        last_col = len(_MATCH_HEADERS)
        ws.auto_filter.ref = f"A1:{ws.cell(row=last_row, column=last_col).coordinate}"

    ws.freeze_panes = "A2"

    widths = [5, 22, 30, 46, 20, 22, 30, 46, 20, 16, 18, 18, 20]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width


def _write_log_sheet(ws: Worksheet, log_entries: list[LogEntry]) -> None:
    for col_idx, header in enumerate(_LOG_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    for row_idx, entry in enumerate(log_entries, start=2):
        ws.cell(row=row_idx, column=1, value=entry.timestamp).font = _BODY_FONT
        ws.cell(row=row_idx, column=2, value=entry.level.value).font = _BODY_FONT
        message_cell = ws.cell(row=row_idx, column=3, value=entry.message)
        message_cell.font = _BODY_FONT
        message_cell.alignment = Alignment(wrap_text=False)

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 100
